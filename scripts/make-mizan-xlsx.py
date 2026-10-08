#!/usr/bin/env python3
"""Build the downloadable Mizan Excel ledger (formulas, RTL, data validation)."""

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from pathlib import Path

OUT = Path("/workspace/public/mizan.xlsx")
OUT_ARTIFACT = Path("/workspace/artifacts/mizan-دفتر-التجارة.xlsx")

INK = "1A1916"
PAPER = "F7F4EE"
CARD = "FFFFFF"
MUTED = "6B6560"
LINE = "DDD6CC"
INCOME = "3F6B45"
EXPENSE = "A35A4A"
CREAM = "E8E2D6"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
fill_paper = PatternFill("solid", fgColor=PAPER)
fill_card = PatternFill("solid", fgColor=CARD)
fill_ink = PatternFill("solid", fgColor=INK)
fill_cream = PatternFill("solid", fgColor=CREAM)
fill_income = PatternFill("solid", fgColor="E7F0E6")
fill_expense = PatternFill("solid", fgColor="F6E8E4")
fill_head = PatternFill("solid", fgColor=INK)

font_title = Font(name="Calibri", size=22, bold=True, color="F0EBE3")
font_h = Font(name="Calibri", size=14, bold=True, color=INK)
font_label = Font(name="Calibri", size=11, color=MUTED)
font_value = Font(name="Calibri", size=16, bold=True, color=INK)
font_body = Font(name="Calibri", size=11, color=INK)
font_head = Font(name="Calibri", size=11, bold=True, color="F0EBE3")
font_income = Font(name="Calibri", size=16, bold=True, color=INCOME)
font_expense = Font(name="Calibri", size=16, bold=True, color=EXPENSE)
font_profit = Font(name="Calibri", size=18, bold=True, color=INK)
align_r = Alignment(horizontal="right", vertical="center", wrap_text=True)
align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)

EXPENSE_CATS = [
    "شراء البضاعة",
    "الإشهار",
    "التوصيل",
    "التغليف",
    "عمولة التوصيل",
    "المرتجعات",
    "رسوم المنصة",
    "تخزين وإيجار",
    "مصروف آخر",
]
INCOME_CATS = ["مبيعات", "استرجاع مصروف", "مدخول آخر"]
ALL_CATS = INCOME_CATS + EXPENSE_CATS

EXAMPLES = [
    ("2026-09-28", "مدخول", "مبيعات", 21400, "مبيعات إنستغرام — 6 طلبات"),
    ("2026-09-28", "مصروف", "التوصيل", 3300, "ياليدين"),
    ("2026-09-27", "مصروف", "الإشهار", 6500, "إشهار فيسبوك"),
    ("2026-09-26", "مدخول", "مبيعات", 17800, "مبيعات تيك توك"),
    ("2026-09-24", "مصروف", "شراء البضاعة", 42000, "شراء سلعة من المورد"),
    ("2026-09-22", "مدخول", "مبيعات", 25600, "مبيعات الويكاند"),
    ("2026-09-20", "مصروف", "عمولة التوصيل", 1800, "كاش أون ديليفري"),
    ("2026-09-18", "مصروف", "التغليف", 2400, "أكياس وكروتونة"),
    ("2026-09-15", "مدخول", "مبيعات", 28400, "حملة الجمعة"),
    ("2026-09-15", "مصروف", "الإشهار", 9100, "رييلز ممولة"),
]


def style_cell(cell, *, font=None, fill=None, align=None, num=None, border=True):
    cell.font = font or font_body
    cell.fill = fill or fill_card
    cell.alignment = align or align_r
    if border:
        cell.border = thin
    if num:
        cell.number_format = num


def set_col_widths(ws, widths):
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w


def build():
    wb = Workbook()
    wb.calculation.calcMode = "auto"

    lists = wb.active
    lists.title = "قوائم"
    dash = wb.create_sheet("لوحة", 0)
    ledger = wb.create_sheet("الحركات", 1)
    how = wb.create_sheet("كيفاش تخدم", 3)

    # --- hidden lists for dropdowns ---
    lists.sheet_view.rightToLeft = True
    lists["A1"] = "النوع"
    lists["B1"] = "الأبواب"
    lists["A2"] = "مدخول"
    lists["A3"] = "مصروف"
    for i, c in enumerate(ALL_CATS, start=2):
        lists.cell(i, 2, c)
    lists.sheet_state = "hidden"

    # --- ledger ---
    ledger.sheet_view.rightToLeft = True
    ledger.sheet_view.showGridLines = False
    ledger.freeze_panes = "A2"
    ledger.row_dimensions[1].height = 28
    headers = ["التاريخ", "النوع", "الباب", "المبلغ (دج)", "ملاحظة"]
    for i, h in enumerate(headers, start=1):
        cell = ledger.cell(1, i, h)
        style_cell(cell, font=font_head, fill=fill_head, align=align_c)

    last_row = 201
    for r in range(2, last_row + 1):
        ledger.row_dimensions[r].height = 22
        date_c = ledger.cell(r, 1)
        type_c = ledger.cell(r, 2)
        cat_c = ledger.cell(r, 3)
        amt_c = ledger.cell(r, 4)
        note_c = ledger.cell(r, 5)
        style_cell(date_c, num="YYYY-MM-DD")
        style_cell(type_c, align=align_c)
        style_cell(cat_c)
        style_cell(amt_c, num="# ##0")
        style_cell(note_c)
        if r <= len(EXAMPLES) + 1:
            ex = EXAMPLES[r - 2]
            date_c.value = ex[0]
            type_c.value = ex[1]
            cat_c.value = ex[2]
            amt_c.value = ex[3]
            note_c.value = ex[4]
            if ex[1] == "مدخول":
                amt_c.font = Font(name="Calibri", size=11, bold=True, color=INCOME)
                type_c.fill = fill_income
            else:
                amt_c.font = Font(name="Calibri", size=11, bold=True, color=EXPENSE)
                type_c.fill = fill_expense

    dv_type = DataValidation(type="list", formula1="قوائم!$A$2:$A$3", allow_blank=True)
    dv_type.error = "اختار مدخول أو مصروف"
    dv_type.errorTitle = "النوع"
    dv_type.prompt = "مدخول = فلوس دخلات · مصروف = فلوس خرجات"
    dv_type.promptTitle = "النوع"
    dv_cat = DataValidation(type="list", formula1="قوائم!$B$2:$B$13", allow_blank=True)
    dv_cat.error = "اختار باب من القائمة"
    dv_cat.errorTitle = "الباب"
    ledger.add_data_validation(dv_type)
    ledger.add_data_validation(dv_cat)
    dv_type.add(f"B2:B{last_row}")
    dv_cat.add(f"C2:C{last_row}")

    table = Table(displayName="Harakat", ref=f"A1:E{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True
    )
    ledger.add_table(table)
    set_col_widths(ledger, {"A": 16, "B": 14, "C": 20, "D": 16, "E": 42})
    ledger.auto_filter.ref = f"A1:E{last_row}"
    ledger.sheet_properties.pageSetUpPr.fitToPage = True

    # conditional color on type
    ledger.conditional_formatting.add(
        f"B2:B{last_row}",
        CellIsRule(operator="equal", formula=['"مدخول"'], fill=fill_income),
    )
    ledger.conditional_formatting.add(
        f"B2:B{last_row}",
        CellIsRule(operator="equal", formula=['"مصروف"'], fill=fill_expense),
    )

    # --- dashboard ---
    dash.sheet_view.rightToLeft = True
    dash.sheet_view.showGridLines = False
    dash.page_setup.fitToPage = True
    dash.sheet_properties.tabColor = "8FAF88"
    for r in range(1, 40):
        dash.row_dimensions[r].height = 22
        for c in range(1, 8):
            dash.cell(r, c).fill = fill_paper
            dash.cell(r, c).border = Border()

    dash.merge_cells("A1:F1")
    dash.row_dimensions[1].height = 46
    title = dash["A1"]
    title.value = "ميزان — دفتر تجارتك الإلكترونية"
    title.font = font_title
    title.fill = fill_ink
    title.alignment = Alignment(horizontal="right", vertical="center")
    for c in range(1, 7):
        dash.cell(1, c).fill = fill_ink

    dash.merge_cells("A2:F2")
    dash.row_dimensions[2].height = 28
    sub = dash["A2"]
    sub.value = "الأرقام تتبدّل وحدها كي تسجّل في ورقة «الحركات». المبالغ بالدينار الجزائري (دج)."
    sub.font = font_label
    sub.fill = fill_paper
    sub.alignment = align_r

    # KPI cards
    def kpi(row, col, label, formula, value_font, fill):
        lab = dash.cell(row, col, label)
        style_cell(lab, font=font_label, fill=fill, align=align_r)
        val = dash.cell(row + 1, col, formula)
        style_cell(val, font=value_font, fill=fill, align=align_r, num='# ##0" دج"')
        dash.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        dash.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        dash.row_dimensions[row].height = 20
        dash.row_dimensions[row + 1].height = 32

    kpi(4, 1, "المداخيل — شحال دخل", '=SUMIF(الحركات!B:B,"مدخول",الحركات!D:D)', font_income, fill_income)
    kpi(4, 3, "المصاريف — شحال خرج", '=SUMIF(الحركات!B:B,"مصروف",الحركات!D:D)', font_expense, fill_expense)
    kpi(4, 5, "الربح الصافي", "=A5-C5", font_profit, fill_cream)

    dash.merge_cells("A7:B7")
    dash["A7"] = "هامش الربح"
    style_cell(dash["A7"], font=font_label, fill=fill_cream)
    dash.merge_cells("A8:B8")
    dash["A8"] = '=IF(A5=0,"—",A5 and 1)'
    dash["A8"].value = '=IF(A5=0,"—",E5/A5)'
    style_cell(dash["A8"], font=font_profit, fill=fill_cream, num="0%")

    dash.merge_cells("C7:D7")
    dash["C7"] = "عائد الإشهار (كل 1 دج إشهار)"
    style_cell(dash["C7"], font=font_label, fill=fill_income)
    dash.merge_cells("C8:D8")
    dash["C8"].value = '=IF(SUMIF(الحركات!C:C,"الإشهار",الحركات!D:D)=0,"—",A5/SUMIF(الحركات!C:C,"الإشهار",الحركات!D:D))'
    style_cell(dash["C8"], font=font_income, fill=fill_income, num='0.0" دج"')

    dash.merge_cells("E7:F7")
    dash["E7"] = "هذا الشهر — المداخيل"
    style_cell(dash["E7"], font=font_label, fill=fill_card)
    dash.merge_cells("E8:F8")
    dash["E8"].value = (
        '=SUMIFS(الحركات!D:D,الحركات!B:B,"مدخول",الحركات!A:A,">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),'
        'الحركات!A:A,"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1))'
    )
    style_cell(dash["E8"], font=font_income, fill=fill_card, num='# ##0" دج"')

    dash.merge_cells("A10:F10")
    dash["A10"] = "وين راحوا المصاريف"
    style_cell(dash["A10"], font=font_h, fill=fill_paper, border=False)
    dash.row_dimensions[10].height = 28

    dash["A11"] = "الباب"
    dash["B11"] = "المبلغ"
    dash["C11"] = "النسبة"
    for col in ("A", "B", "C"):
        style_cell(dash[f"{col}11"], font=font_head, fill=fill_head, align=align_c)

    for i, cat in enumerate(EXPENSE_CATS):
        r = 12 + i
        dash.cell(r, 1, cat)
        style_cell(dash.cell(r, 1), fill=fill_card)
        dash.cell(r, 2, f'=SUMIF(الحركات!C:C,A{r},الحركات!D:D)')
        style_cell(dash.cell(r, 2), fill=fill_card, num='# ##0" دج"')
        dash.cell(r, 3, f'=IF(C5=0,0,B{r}/C5)')
        style_cell(dash.cell(r, 3), fill=fill_card, num="0%")

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "المصاريف حسب الباب"
    chart.y_axis.title = None
    chart.x_axis.title = None
    data = Reference(dash, min_col=2, min_row=11, max_row=20)
    cats = Reference(dash, min_col=1, min_row=12, max_row=20)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.width = 18
    chart.height = 8
    dash.add_chart(chart, "E11")

    dash.merge_cells("A22:F22")
    dash["A22"] = (
        "باش تبدى بحساباتك: روح لورقة «الحركات»، امسح صفوف المثال، وسكّل مداخيل ومصاريف كل يوم. "
        "اللوحة تحسب وحدها. تقدر تصفي الجدول من السهم فوق العمود."
    )
    style_cell(dash["A22"], font=font_label, fill=fill_paper, border=False)
    dash.row_dimensions[22].height = 40

    set_col_widths(dash, {"A": 26, "B": 16, "C": 26, "D": 16, "E": 22, "F": 18})

    # --- how to ---
    how.sheet_view.rightToLeft = True
    how.sheet_view.showGridLines = False
    for r in range(1, 24):
        how.row_dimensions[r].height = 22
        for c in range(1, 4):
            how.cell(r, c).fill = fill_paper
    how.merge_cells("A1:C1")
    how.row_dimensions[1].height = 40
    how["A1"] = "كيفاش تخدم دفتر ميزان"
    how["A1"].font = font_title
    how["A1"].fill = fill_ink
    how["A1"].alignment = Alignment(horizontal="right", vertical="center")
    for c in range(1, 4):
        how.cell(1, c).fill = fill_ink

    steps = [
        "1) ورقة «الحركات» هي الدفتر. كل سطر = حركة واحدة (مدخول أو مصروف).",
        "2) اختار النوع من القائمة: مدخول (دخل للخزنة) أو مصروف (خرج من اليد).",
        "3) اختار الباب: مبيعات، إشهار، توصيل، شراء البضاعة…",
        "4) اكتب المبلغ بالدينار، بلا فاصلة. مثال: 4500",
        "5) التاريخ بصيغة 2026-10-01. الملاحظة اختيارية (ياليدين، حملة رييلز…).",
        "6) ارجع للوحة: المداخيل، المصاريف، الربح، هامش الربح وعائد الإشهار يتحسبو وحدهم.",
        "7) امسح صفوف المثال قبل ما تبدأ بحساباتك الحقيقية.",
        "8) تصفية: اضغط السهم فوق عمود الباب أو النوع باش تشوف غير الإشهار مثلا.",
        "9) هذا الملف يتخدم في Excel، Google Sheets، و Numbers. فعّل التعديل إذا طلب منك.",
        "10) التطبيق ميزان على الويب هو نفس الحساب، وفيه تصدير CSV من القائمة.",
    ]
    for i, line in enumerate(steps, start=3):
        how.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        how.row_dimensions[i].height = 28
        cell = how.cell(i, 1, line)
        cell.font = font_body
        cell.fill = fill_paper
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    how.merge_cells("A15:C16")
    how["A15"] = (
        "نصيحة للتجارة الإلكترونية في الجزائر: سجّل المبيعات مجمّعة لليوم (ماشي كل طلب وحدو إذا ما تحبش)، "
        "والمصروف الكبير (سلعة، إشهار، فاتورة ياليدين) كي يخرج من يدك. الربح = الدخل − الخروج، ماشي رقم المبيعات برك."
    )
    how["A15"].font = font_label
    how["A15"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    how["A15"].fill = fill_cream
    set_col_widths(how, {"A": 40, "B": 24, "C": 24})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT_ARTIFACT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    wb.save(OUT_ARTIFACT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
    print(f"wrote {OUT_ARTIFACT} ({OUT_ARTIFACT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
